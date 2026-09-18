"""Relatório da conciliação em planilha: resumo, avisos e pendências."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.domain.exportacao.formatos import _celula_segura
from src.schemas.conciliacao_bancaria import RelatorioConciliacao

ROTULOS = {
    "CONCILIADO": "Conciliado",
    "DATA_DIFERENTE": "Conciliado com data diferente",
    "AGRUPADO": "Agrupado",
    "LOTE_SISPAG": "Lote do SISPAG",
    "LOTE_SISPAG_DIVERGENTE": "Lote do SISPAG com pagamento fora do razão",
    "LOTE_DO_DIA": "Lote do dia (sobra do dia fecha)",
    "DUPLICIDADE_RAZAO": "Possível duplicidade no razão",
    "DUPLICIDADE_EXTRATO": "Possível duplicidade no extrato",
    "VALOR_DIVERGENTE": "Valor divergente",
    "SO_RAZAO": "Só no razão",
    "SO_EXTRATO": "Só no extrato",
}
_NEGRITO = Font(bold=True)
_CABECALHO = PatternFill("solid", fgColor="DDEBF7")


def _linha(ws, valores, negrito=False):
    ws.append([_celula_segura(v) for v in valores])
    if negrito:
        for c in ws[ws.max_row]:
            c.font = _NEGRITO


def gerar_planilha(relatorio: RelatorioConciliacao) -> bytes:
    r = relatorio.resumo
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumo"
    for rotulo, valor in [
        ("Empresa", r.empresa),
        ("Conta do razão", r.conta_razao),
        ("Conta bancária", r.conta_bancaria),
        ("Período", f"{r.periodo_inicio:%d/%m/%Y} a {r.periodo_fim:%d/%m/%Y}"),
        ("Lançamentos no razão", r.lancamentos_razao),
        ("Lançamentos no extrato", r.lancamentos_extrato),
        ("Conciliados", r.conciliados),
        ("Pendências", r.pendencias),
        ("Aplicação automática não conferida (extrato não traz)", r.aplicacao_sem_extrato),
        ("Movimento do razão (sem abertura)", float(r.movimento_razao)),
        ("Movimento do extrato", float(r.movimento_extrato)),
        ("Diferença", float(r.diferenca)),
        ("As pendências explicam a diferença?", "Sim" if r.diferenca_explicada else "NÃO — revisar"),
    ]:
        _linha(ws, [rotulo, valor])
        ws.cell(ws.max_row, 1).font = _NEGRITO
        if isinstance(valor, float):
            ws.cell(ws.max_row, 2).number_format = "#,##0.00"
    for a in r.abertura:
        _linha(ws, ["Abertura (não conciliável)", f"{a.data:%d/%m/%Y} {a.historico} {a.valor:,.2f}"])
    if relatorio.avisos:
        ws.append([])
        _linha(ws, ["Avisos"], negrito=True)
        for aviso in relatorio.avisos:
            _linha(ws, [aviso])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 90

    wp = wb.create_sheet("Pendências")
    colunas = ["Tipo", "Data razão", "Valor razão", "Histórico razão", "Lote", "Contrapartida",
               "Data extrato", "Valor extrato", "Histórico extrato", "Diferença"]
    _linha(wp, colunas, negrito=True)
    for c in wp[1]:
        c.fill = _CABECALHO
    for g in relatorio.pendencias:
        altura = max(len(g.razao), len(g.extrato), 1)
        for i in range(altura):
            rz = g.razao[i] if i < len(g.razao) else None
            ex = g.extrato[i] if i < len(g.extrato) else None
            _linha(wp, [
                ROTULOS.get(g.tipo, g.tipo) if i == 0 else "",
                rz.data if rz else None, float(rz.valor) if rz else None,
                rz.historico if rz else None, rz.lote if rz else None, rz.contrapartida if rz else None,
                ex.data if ex else None, float(ex.valor) if ex else None,
                ex.historico if ex else None,
                float(g.diferenca) if i == 0 else None,
            ])
            for col in (2, 7):
                wp.cell(wp.max_row, col).number_format = "DD/MM/YYYY"
            for col in (3, 8, 10):
                wp.cell(wp.max_row, col).number_format = "#,##0.00"
        for pg in g.sispag_faltando:
            # Pago pelo banco dentro do lote, sem lançamento no razão.
            _linha(wp, [
                "  pagamento do SISPAG sem par no razão", None, None,
                f"{pg.favorecido} ({pg.documento}) — {pg.tipo}", None, None,
                pg.data, float(pg.valor), None, None,
            ])
            wp.cell(wp.max_row, 7).number_format = "DD/MM/YYYY"
            wp.cell(wp.max_row, 8).number_format = "#,##0.00"
    for i, w in enumerate([30, 12, 15, 55, 12, 14, 12, 15, 45, 14], 1):
        wp.column_dimensions[get_column_letter(i)].width = w
    wp.freeze_panes = "A2"
    wp.auto_filter.ref = f"A1:J{wp.max_row}"

    if relatorio.por_dia:
        wd = wb.create_sheet("Por dia")
        _linha(wd, ["Data", "Lanç. razão", "Lanç. extrato", "Movimento razão", "Movimento extrato",
                    "Diferença", "Pendências", "Casados em outro dia", "Aplicação não conferida"],
               negrito=True)
        for c in wd[1]:
            c.fill = _CABECALHO
        for d in relatorio.por_dia:
            _linha(wd, [d.data, d.lancamentos_razao, d.lancamentos_extrato,
                        float(d.movimento_razao), float(d.movimento_extrato), float(d.diferenca),
                        d.pendencias, d.data_diferente, d.aplicacao_sem_extrato])
            wd.cell(wd.max_row, 1).number_format = "DD/MM/YYYY"
            for col in (4, 5, 6):
                wd.cell(wd.max_row, col).number_format = "#,##0.00"
        for i, w in enumerate([12, 12, 13, 17, 17, 15, 12, 18, 22], 1):
            wd.column_dimensions[get_column_letter(i)].width = w
        wd.freeze_panes = "A2"
        wd.auto_filter.ref = f"A1:I{wd.max_row}"

    if relatorio.aplicacao_sem_extrato:
        wa = wb.create_sheet("Aplicação não conferida")
        _linha(wa, ["Data", "Valor", "Histórico", "Lote", "Contrapartida"], negrito=True)
        for c in wa[1]:
            c.fill = _CABECALHO
        for linha in relatorio.aplicacao_sem_extrato:
            _linha(wa, [linha.data, float(linha.valor), linha.historico,
                        linha.lote, linha.contrapartida])
            wa.cell(wa.max_row, 1).number_format = "DD/MM/YYYY"
            wa.cell(wa.max_row, 2).number_format = "#,##0.00"
        for i, w in enumerate([12, 15, 55, 12, 14], 1):
            wa.column_dimensions[get_column_letter(i)].width = w
        wa.freeze_panes = "A2"

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
