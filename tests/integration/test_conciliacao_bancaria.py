"""Conciliação bancária: razão enviado × extrato importado, pela rota.

O razão é sintético, no layout do relatório RAZÃO do Mister Contador — o mesmo
que o ConcilPro lê para fornecedores, com um bloco só.
"""

from __future__ import annotations

import io
import uuid
from datetime import datetime

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from src.db.models import AgenciaBancaria, Empresa, PlanoConta, Tenant, Transacao, Usuario

_OFX = """\
OFXHEADER:100
DATA:OFXSGML
VERSION:102
<OFX><BANKMSGSRSV1><STMTTRNRS><STMTRS><BANKTRANLIST>
<STMTTRN><TRNTYPE>CREDIT<DTPOSTED>20250203<TRNAMT>1200.00<FITID>CB1<MEMO>PIX RECEBIDO CLIENTE ALFA</STMTTRN>
<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20250210<TRNAMT>-350.00<FITID>CB2<MEMO>SISPAG ENERGIA</STMTTRN>
<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20250226<TRNAMT>-90.00<FITID>CB3<MEMO>TARIFA BANCARIA</STMTTRN>
</BANKTRANLIST></STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>
"""


def _razao_xlsx(lancamentos, *, cnpj="12.345.678/0001-95", conta=(2699, "1.1.1.02.0015", "BANCO ITAU C/C 122870"),
                total=None, contas_extras=0) -> bytes:
    """lancamentos: (dia, lote, historico, contra, debito, credito)."""
    wb = openpyxl.Workbook()
    ws = wb.active

    def linha(**cols):
        valores = [None] * 15
        for i, v in cols.items():
            valores[int(i[1:])] = v
        ws.append(valores)

    linha(c0="Empresa:", c2="DECATEC LTDA", c12="Folha:", c14=1)
    linha(c0="C.N.P.J.:", c2=cnpj)
    linha(c0="Período:", c2="01/02/2025 - 28/02/2025")
    linha(c0="RAZÃO")
    linha(c0="Data", c1="Lote", c2="Histórico", c8="Cta.C.Part.", c9="Débito", c10="Crédito", c13="Saldo-Exercício")
    for bloco in range(1 + contas_extras):
        codigo, classif, nome = conta if bloco == 0 else (9000 + bloco, f"1.1.1.02.9{bloco:03d}", "OUTRO BANCO")
        linha(c0="Conta:", c1=codigo, c2=classif, c5=nome)
        linha(c2="SALDO ANTERIOR", c13=0.0)
        saldo, deb, cred = 0.0, 0.0, 0.0
        for dia, lote, hist, contra, d, c in lancamentos:
            saldo += (d or 0) - (c or 0)
            deb += d or 0
            cred += c or 0
            linha(c0=datetime(2025, 2, dia), c1=lote, c2=hist, c8=contra, c9=d, c10=c, c13=saldo)
        td, tc = total or (deb, cred)
        linha(c0="Total da conta:", c9=td, c10=tc)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


_RAZAO_PADRAO = [
    (1, 3026001, "RETORNO DE SALDO NEGATIVO - ITAU", 1370, None, 500.00),
    (3, 3137001, "RECEBIMENTO CLIENTE ALFA", 504, 1200.00, None),
    (10, 3137002, "PAGTO ENERGIA ELETRICA - COPEL", 354, None, 350.00),
    (10, 3026002, "TRANSFERENCIA ENTRE CONTAS", 554, None, 350.00),   # duplicidade
    (20, 3137003, "PGTO FORNECEDOR SEM PAR", 506, None, 42.00),        # só no razão
]  # a tarifa de 90,00 do dia 26 fica só no extrato


async def _login(client: AsyncClient, tenant: Tenant, usuario: Usuario) -> str:
    r = await client.post(
        "/api/v1/auth/login",
        json={"tenant_id": str(tenant.id), "email": usuario.email, "senha": "senha_segura_123"},
    )
    assert r.status_code == 200
    return r.json()["csrf_token"]


async def _conta_com_extrato(client, empresa, csrf) -> dict:
    agencia = (await client.post(
        f"/api/v1/empresas/{empresa.id}/agencias",
        json={"banco_sigla": "ITAU", "agencia": "7285", "numero": "12287"},
        headers={"X-CSRF-Token": csrf},
    )).json()
    r = await client.post(
        f"/api/v1/empresas/{empresa.id}/extrato/importar?agencia_id={agencia['id']}",
        files={"arquivo": ("fev.ofx", io.BytesIO(_OFX.encode()), "application/octet-stream")},
        headers={"X-CSRF-Token": csrf},
    )
    assert r.status_code == 202, r.text
    return agencia


async def _conciliar(client, empresa, agencia_id, csrf, razao: bytes, formato="json", nome="razao.xlsx"):
    return await client.post(
        f"/api/v1/empresas/{empresa.id}/concilpro/razao-extrato"
        f"?agencia_id={agencia_id}&formato={formato}",
        files={"arquivo": (nome, io.BytesIO(razao), "application/octet-stream")},
        headers={"X-CSRF-Token": csrf},
    )


@pytest.mark.asyncio
async def test_relatorio_classifica_as_pendencias_e_explica_a_diferenca(
    client: AsyncClient, tenant: Tenant, usuario: Usuario, empresa: Empresa
):
    csrf = await _login(client, tenant, usuario)
    agencia = await _conta_com_extrato(client, empresa, csrf)

    r = await _conciliar(client, empresa, agencia["id"], csrf, _razao_xlsx(_RAZAO_PADRAO))

    assert r.status_code == 200, r.text
    corpo = r.json()
    resumo = corpo["resumo"]
    assert resumo["conciliados"] == 2
    assert [a["historico"] for a in resumo["abertura"]] == ["RETORNO DE SALDO NEGATIVO - ITAU"]
    assert resumo["diferenca_explicada"] is True

    por_tipo = {g["tipo"]: g for g in corpo["pendencias"]}
    assert set(por_tipo) == {"DUPLICIDADE_RAZAO", "SO_RAZAO", "SO_EXTRATO"}
    # O histórico decidiu qual dos dois 350,00 é o par.
    assert por_tipo["DUPLICIDADE_RAZAO"]["razao"][0]["historico"] == "TRANSFERENCIA ENTRE CONTAS"
    assert por_tipo["DUPLICIDADE_RAZAO"]["razao"][0]["lote"] == "3026002"
    assert por_tipo["SO_RAZAO"]["razao"][0]["historico"] == "PGTO FORNECEDOR SEM PAR"
    assert por_tipo["SO_EXTRATO"]["extrato"][0]["historico"] == "TARIFA BANCARIA"


@pytest.mark.asyncio
async def test_conciliar_nao_grava_nada(
    client: AsyncClient, db: AsyncSession, tenant: Tenant, usuario: Usuario, empresa: Empresa
):
    """Só relata: nenhuma transação é criada, removida ou alterada."""
    csrf = await _login(client, tenant, usuario)
    agencia = await _conta_com_extrato(client, empresa, csrf)
    antes = [(t.id, t.status, t.deleted_at) for t in (await db.execute(select(Transacao))).scalars()]

    await _conciliar(client, empresa, agencia["id"], csrf, _razao_xlsx(_RAZAO_PADRAO))

    depois = [(t.id, t.status, t.deleted_at) for t in (await db.execute(select(Transacao))).scalars()]
    assert antes == depois


@pytest.mark.asyncio
async def test_exporta_em_planilha(client: AsyncClient, tenant: Tenant, usuario: Usuario, empresa: Empresa):
    csrf = await _login(client, tenant, usuario)
    agencia = await _conta_com_extrato(client, empresa, csrf)

    r = await _conciliar(client, empresa, agencia["id"], csrf, _razao_xlsx(_RAZAO_PADRAO), formato="xlsx")

    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Resumo", "Pendências"]
    tipos = [row[0] for row in wb["Pendências"].iter_rows(min_row=2, values_only=True) if row[0]]
    assert "Possível duplicidade no razão" in tipos


@pytest.mark.asyncio
async def test_razao_de_outra_empresa_e_recusado(
    client: AsyncClient, tenant: Tenant, usuario: Usuario, empresa: Empresa
):
    csrf = await _login(client, tenant, usuario)
    agencia = await _conta_com_extrato(client, empresa, csrf)

    r = await _conciliar(client, empresa, agencia["id"], csrf,
                         _razao_xlsx(_RAZAO_PADRAO, cnpj="17.122.471/0001-75"))

    assert r.status_code == 422
    assert "outra empresa" in r.json()["message"]


@pytest.mark.asyncio
async def test_razao_de_outra_conta_contabil_e_recusado(
    client: AsyncClient, db: AsyncSession, tenant: Tenant, usuario: Usuario, empresa: Empresa
):
    """Razão de um banco com o extrato de outro: pendências falsas do começo ao fim."""
    csrf = await _login(client, tenant, usuario)
    agencia = await _conta_com_extrato(client, empresa, csrf)
    conta = PlanoConta(empresa_id=empresa.id, codigo="1.1.1.02.0020", descricao="BANCO BRADESCO",
                       tipo="ativo", conta_numero=3100)
    db.add(conta)
    await db.flush()
    ag = await db.get(AgenciaBancaria, uuid.UUID(agencia["id"]))
    ag.conta_contabil_id = conta.id
    await db.flush()

    r = await _conciliar(client, empresa, agencia["id"], csrf, _razao_xlsx(_RAZAO_PADRAO))

    assert r.status_code == 422
    assert "2699" in r.json()["message"] and "BANCO BRADESCO" in r.json()["message"]


@pytest.mark.asyncio
async def test_conta_contabil_vinculada_certa_passa_sem_aviso(
    client: AsyncClient, db: AsyncSession, tenant: Tenant, usuario: Usuario, empresa: Empresa
):
    csrf = await _login(client, tenant, usuario)
    agencia = await _conta_com_extrato(client, empresa, csrf)
    conta = PlanoConta(empresa_id=empresa.id, codigo="1.1.1.02.0015", descricao="BANCO ITAU",
                       tipo="ativo", conta_numero=2699)
    db.add(conta)
    await db.flush()
    ag = await db.get(AgenciaBancaria, uuid.UUID(agencia["id"]))
    ag.conta_contabil_id = conta.id
    await db.flush()

    r = await _conciliar(client, empresa, agencia["id"], csrf, _razao_xlsx(_RAZAO_PADRAO))

    assert r.status_code == 200, r.text
    assert not any("vinculada" in a for a in r.json()["avisos"])


@pytest.mark.asyncio
async def test_sem_extrato_importado_no_periodo_e_recusado(
    client: AsyncClient, tenant: Tenant, usuario: Usuario, empresa: Empresa
):
    csrf = await _login(client, tenant, usuario)
    agencia = (await client.post(
        f"/api/v1/empresas/{empresa.id}/agencias",
        json={"banco_sigla": "ITAU", "agencia": "7285", "numero": "12287"},
        headers={"X-CSRF-Token": csrf},
    )).json()

    r = await _conciliar(client, empresa, agencia["id"], csrf, _razao_xlsx(_RAZAO_PADRAO))

    assert r.status_code == 422
    assert "Importe o extrato" in r.json()["message"]


@pytest.mark.asyncio
async def test_razao_com_varias_contas_e_recusado(
    client: AsyncClient, tenant: Tenant, usuario: Usuario, empresa: Empresa
):
    csrf = await _login(client, tenant, usuario)
    agencia = await _conta_com_extrato(client, empresa, csrf)

    r = await _conciliar(client, empresa, agencia["id"], csrf, _razao_xlsx(_RAZAO_PADRAO, contas_extras=1))

    assert r.status_code == 422
    assert "uma conta bancária por vez" in r.json()["message"]


@pytest.mark.asyncio
async def test_razao_que_nao_fecha_com_o_total_e_recusado(
    client: AsyncClient, tenant: Tenant, usuario: Usuario, empresa: Empresa
):
    """Leitura pela metade apontaria pendências que não existem."""
    csrf = await _login(client, tenant, usuario)
    agencia = await _conta_com_extrato(client, empresa, csrf)

    r = await _conciliar(client, empresa, agencia["id"], csrf,
                         _razao_xlsx(_RAZAO_PADRAO, total=(1200.00, 9999.00)))

    assert r.status_code == 422
    assert "Total da conta" in r.json()["message"]


@pytest.mark.asyncio
async def test_sem_csrf_rejeita(client: AsyncClient, tenant: Tenant, usuario: Usuario, empresa: Empresa):
    await _login(client, tenant, usuario)
    r = await client.post(
        f"/api/v1/empresas/{empresa.id}/concilpro/razao-extrato?agencia_id={uuid.uuid4()}",
        files={"arquivo": ("r.xlsx", io.BytesIO(b"PK"), "application/octet-stream")},
    )
    assert r.status_code == 403
